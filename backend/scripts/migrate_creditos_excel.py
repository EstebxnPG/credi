from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app.core.config import settings  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.db.session import SessionLocal  # noqa: E402


MIN_DATE = date(2018, 1, 1)
MAX_DATE = date(2026, 12, 31)
TODAY = date(2026, 7, 2)
SHEET_NAME = "2018-2026"
SYSTEM_USER_NAME = "CREDICONFIEMOS"
SYSTEM_USER_EMAIL = "crediconfiemos@crediconfiemos.com"
SYSTEM_USER_DOCUMENT = "9999999999"


@dataclass
class PreparedData:
    rows: list[dict[str, Any]]
    migratable: list[dict[str, Any]]
    rejected: list[dict[str, Any]]
    pensionados: list[dict[str, Any]]
    duplicate_libranzas: set[str]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def normalize_text(value: Any) -> str:
    text_value = clean_text(value)
    if not text_value:
        return ""
    normalized = unicodedata.normalize("NFKD", text_value)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.upper()
    normalized = re.sub(r"[^A-ZÑ0-9 ]+", " ", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def clean_document(value: Any) -> str:
    digits = re.sub(r"\D", "", clean_text(value))
    if len(digits) < 5 or set(digits) == {"0"}:
        return ""
    return digits


def clean_libranza(value: Any) -> str:
    value = clean_text(value).upper()
    if not value or value in {"0", "?", "NO.LIB", "NAN"}:
        return ""
    return value[:50]


def clean_phone(value: Any) -> str:
    digits = re.sub(r"\D", "", clean_text(value))
    if len(digits) < 7:
        return ""
    return digits[:20]


def clean_address(value: Any) -> str:
    value = clean_text(value)
    if not value:
        return ""
    if normalize_text(value) in {"CANCELO", "FALLECIO", "NEGADO", "BUZON", "CORREO VOZ"}:
        return ""
    return value[:200]


def parse_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    digits = re.sub(r"[$\s,.]", "", clean_text(value))
    return int(digits) if re.fullmatch(r"\d+", digits) else None


def normalize_date_text(value: Any) -> str:
    original = clean_text(value)
    fixed = re.sub(r"([/-])2(20\d{2})$", r"\1\2", original)
    fixed = re.sub(r"([/-])(20\d{2})\d$", r"\1\2", fixed)
    fixed = re.sub(r"^(\d{2})(\d{2})/(20\d{2})$", r"\1/\2/\3", fixed)
    return fixed


def parse_date(value: Any) -> tuple[date | None, str]:
    if value is None or clean_text(value) == "":
        return None, "FALTA_FECHA"
    if isinstance(value, datetime):
        parsed = value.date()
        return (parsed, "OK") if MIN_DATE <= parsed <= MAX_DATE else (None, "FUERA_RANGO")
    if isinstance(value, date):
        return (value, "OK") if MIN_DATE <= value <= MAX_DATE else (None, "FUERA_RANGO")
    if isinstance(value, (int, float)) and 20000 <= value <= 60000:
        parsed = (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        return (parsed, "OK_SERIAL_EXCEL") if MIN_DATE <= parsed <= MAX_DATE else (None, "FUERA_RANGO")

    original = clean_text(value)
    fixed = normalize_date_text(original)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(fixed, fmt).date()
            if MIN_DATE <= parsed <= MAX_DATE:
                return parsed, "CORREGIDA" if fixed != original else "OK"
            return None, "FUERA_RANGO"
        except ValueError:
            continue
    return None, "INVALIDA"


def add_months(base: date, months: int) -> date:
    month_index = base.month - 1 + months
    year = base.year + month_index // 12
    month = month_index % 12 + 1
    days = [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return date(year, month, min(base.day, days[month - 1]))


def read_excel(input_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(i for i, row in enumerate(rows) if any(cell is not None for cell in row))
    headers = [clean_text(value) if clean_text(value) else f"extra_col_{idx + 1}" for idx, value in enumerate(rows[header_index])]

    output: list[dict[str, Any]] = []
    for excel_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        data = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}
        if not any(value is not None and clean_text(value) for value in data.values()):
            continue
        tipo = normalize_text(data.get("TIPO CREDITO"))
        fecha = normalize_text(data.get("FECHA_RAD"))
        if tipo in {"TIPO CREDITO", "TIPO"} or fecha == "FECHA RAD":
            continue
        data["fila_excel"] = excel_row
        output.append(data)
    return output


def prepare(input_path: Path) -> PreparedData:
    rows = read_excel(input_path)
    for row in rows:
        parsed, status = parse_date(row.get("FECHA_RAD"))
        row["fecha_parseada_inicial"] = parsed
        row["estado_fecha_inicial"] = status
        row["cc_limpia"] = clean_document(row.get("Cedula"))
        row["nombre_limpio"] = clean_text(row.get("APELLIDOS Y NOMBRES DEL CLIENTE"))
        row["nombre_comparable"] = normalize_text(row.get("APELLIDOS Y NOMBRES DEL CLIENTE"))
        row["libranza_limpia"] = clean_libranza(row.get("No.Lib"))
        row["telefono_limpio"] = clean_phone(row.get("TELEFONO"))
        row["celular_limpio"] = clean_phone(row.get("CELULAR"))
        row["direccion_limpia"] = clean_address(row.get("DIRECCION"))
        row["monto_parseado"] = parse_int(row.get("MONTO"))
        row["plazo_parseado"] = parse_int(row.get("MESES"))
        row["tipo_normalizado"] = normalize_text(row.get("TIPO CREDITO"))
        row["pagaduria_normalizada"] = normalize_text(row.get("PAGADURIA")) or "SIN REGISTRAR"
        row["cooperativa_normalizada"] = normalize_text(row.get("COOPERATIVA")) or "SIN REGISTRAR"

    previous_date: date | None = None
    for row in rows:
        row["fecha_anterior_valida"] = previous_date
        if row["fecha_parseada_inicial"]:
            previous_date = row["fecha_parseada_inicial"]

    next_date: date | None = None
    for row in reversed(rows):
        row["fecha_siguiente_valida"] = next_date
        if row["fecha_parseada_inicial"]:
            next_date = row["fecha_parseada_inicial"]

    for row in rows:
        if row["fecha_parseada_inicial"]:
            row["fecha_final"] = row["fecha_parseada_inicial"]
            row["criterio_fecha"] = "ORIGINAL_O_CORREGIDA"
        elif row["fecha_anterior_valida"] and row["fecha_siguiente_valida"] and row["fecha_anterior_valida"] <= row["fecha_siguiente_valida"]:
            row["fecha_final"] = row["fecha_anterior_valida"]
            row["criterio_fecha"] = "DEDUCIDA_POR_ORDEN"
        elif row["fecha_anterior_valida"] and not row["fecha_siguiente_valida"]:
            row["fecha_final"] = row["fecha_anterior_valida"]
            row["criterio_fecha"] = "DEDUCIDA_SIN_SIGUIENTE"
        elif row["fecha_siguiente_valida"] and not row["fecha_anterior_valida"]:
            row["fecha_final"] = row["fecha_siguiente_valida"]
            row["criterio_fecha"] = "DEDUCIDA_SIN_PREVIA"
        else:
            row["fecha_final"] = None
            row["criterio_fecha"] = "NO_DEDUCIBLE"

    libranza_counts = Counter(row["libranza_limpia"] for row in rows if row["libranza_limpia"])
    duplicate_libranzas = {key for key, count in libranza_counts.items() if count > 1}

    migratable: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    for row in rows:
        errors: list[str] = []
        if not row["fecha_final"]:
            errors.append("FECHA_NO_DEDUCIBLE")
        if not row["cc_limpia"]:
            errors.append("CC_INVALIDA_O_VACIA")
        if not row["nombre_limpio"]:
            errors.append("NOMBRE_VACIO")
        if row["monto_parseado"] is None:
            errors.append("MONTO_INVALIDO_O_VACIO")
        if row["plazo_parseado"] is None:
            errors.append("PLAZO_INVALIDO_O_VACIO")
        if row["libranza_limpia"] in duplicate_libranzas:
            errors.append("LIBRANZA_DUPLICADA")
        row["errores"] = "; ".join(errors)
        if errors:
            rejected.append(row)
            continue
        row["fecha_fin_estimada"] = add_months(row["fecha_final"], row["plazo_parseado"])
        row["estado_credito"] = "Finalizado" if row["fecha_fin_estimada"] < TODAY else "Aprobado"
        migratable.append(row)

    names_by_cc: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in migratable:
        variants = names_by_cc[row["cc_limpia"]]
        if row["nombre_comparable"] and all(item["nombre_comparable"] != row["nombre_comparable"] for item in variants):
            variants.append(
                {
                    "nombre": row["nombre_limpio"],
                    "nombre_comparable": row["nombre_comparable"],
                    "fila_excel": row["fila_excel"],
                }
            )

    pensionados: list[dict[str, Any]] = []
    cc_to_pensionado: dict[str, dict[str, Any]] = {}
    for cc, group_rows in _group_migratable_by_cc(migratable).items():
        variants = names_by_cc[cc]
        chosen = max(variants, key=lambda item: (len(item["nombre_comparable"]), len(item["nombre"]), -item["fila_excel"]))
        latest_phone = latest_non_empty(group_rows, "telefono_limpio")
        latest_mobile = latest_non_empty(group_rows, "celular_limpio")
        latest_address = latest_non_empty(group_rows, "direccion_limpia")
        pensionado = {
            "pensionado_id_simulado": len(pensionados) + 1,
            "documento": cc,
            "nombre": chosen["nombre"][:150],
            "fila_excel_nombre_elegido": chosen["fila_excel"],
            "telefono": latest_phone["value"] if latest_phone else "",
            "fila_excel_telefono": latest_phone["fila_excel"] if latest_phone else "",
            "celular": latest_mobile["value"] if latest_mobile else "",
            "fila_excel_celular": latest_mobile["fila_excel"] if latest_mobile else "",
            "direccion": latest_address["value"] if latest_address else "",
            "fila_excel_direccion": latest_address["fila_excel"] if latest_address else "",
            "cantidad_creditos": len(group_rows),
            "cantidad_nombres_distintos": len(variants),
            "variantes_nombre": " | ".join(item["nombre"] for item in variants),
        }
        pensionados.append(pensionado)
        cc_to_pensionado[cc] = pensionado

    for row in rows:
        pensionado = cc_to_pensionado.get(row["cc_limpia"])
        row["pensionado_id_simulado"] = pensionado["pensionado_id_simulado"] if pensionado else ""
        row["pensionado_nombre_final"] = pensionado["nombre"] if pensionado else ""
        row["fila_excel_nombre_elegido"] = pensionado["fila_excel_nombre_elegido"] if pensionado else ""
        row["telefono_final"] = pensionado["telefono"] if pensionado else ""
        row["fila_excel_telefono"] = pensionado["fila_excel_telefono"] if pensionado else ""
        row["celular_final"] = pensionado["celular"] if pensionado else ""
        row["fila_excel_celular"] = pensionado["fila_excel_celular"] if pensionado else ""
        row["direccion_final"] = pensionado["direccion"] if pensionado else ""
        row["fila_excel_direccion"] = pensionado["fila_excel_direccion"] if pensionado else ""
        row["decision"] = "MIGRAR" if row in migratable else "OMITIR"

    return PreparedData(rows=rows, migratable=migratable, rejected=rejected, pensionados=pensionados, duplicate_libranzas=duplicate_libranzas)


def _group_migratable_by_cc(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in sorted(rows, key=lambda item: item["fila_excel"]):
        grouped.setdefault(row["cc_limpia"], []).append(row)
    return grouped


def latest_non_empty(rows: list[dict[str, Any]], key: str) -> dict[str, Any] | None:
    for row in sorted(rows, key=lambda item: item["fila_excel"], reverse=True):
        if row.get(key):
            return {"value": row[key], "fila_excel": row["fila_excel"]}
    return None


def write_reports(prepared: PreparedData, output_dir: Path) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "dry_run_resumen.json"
    rejected_path = output_dir / "creditos_omitidos.csv"
    workbook_path = output_dir / "dry_run_migracion_creditos.xlsx"

    error_counts = Counter()
    for row in prepared.rejected:
        for error in row["errores"].split("; "):
            if error:
                error_counts[error] += 1

    summary = {
        "filas_analizadas": len(prepared.rows),
        "creditos_migrables": len(prepared.migratable),
        "creditos_omitidos": len(prepared.rejected),
        "pensionados_simulados": len(prepared.pensionados),
        "libranzas_duplicadas_unicas": len(prepared.duplicate_libranzas),
        "errores": dict(error_counts),
    }
    json_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    with rejected_path.open("w", encoding="utf-8-sig", newline="") as file:
        fieldnames = ["fila_excel", "errores", "FECHA_RAD", "Cedula", "cc_limpia", "APELLIDOS Y NOMBRES DEL CLIENTE", "No.Lib", "libranza_limpia", "MONTO", "MESES", "PAGADURIA", "COOPERATIVA"]
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for row in prepared.rejected:
            writer.writerow({field: row.get(field, "") for field in fieldnames})

    write_workbook(prepared, workbook_path, summary, error_counts)
    return {"summary": str(json_path), "rejected": str(rejected_path), "workbook": str(workbook_path)}


def write_workbook(prepared: PreparedData, workbook_path: Path, summary: dict[str, Any], error_counts: Counter) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    def add_sheet(name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        style_sheet(sheet)

    add_sheet("00_resumen", ["metrica", "valor"], [{"metrica": key, "valor": value} for key, value in summary.items() if key != "errores"])
    add_sheet("01_errores", ["error", "filas"], [{"error": key, "filas": value} for key, value in error_counts.most_common()])
    add_sheet(
        "02_pensionados",
        [
            "pensionado_id_simulado",
            "documento",
            "nombre",
            "fila_excel_nombre_elegido",
            "telefono",
            "fila_excel_telefono",
            "celular",
            "fila_excel_celular",
            "direccion",
            "fila_excel_direccion",
            "cantidad_creditos",
            "cantidad_nombres_distintos",
            "variantes_nombre",
        ],
        prepared.pensionados,
    )
    row_headers = [
        "fila_excel", "decision", "errores", "pensionado_id_simulado", "pensionado_nombre_final", "fila_excel_nombre_elegido",
        "telefono_final", "fila_excel_telefono", "celular_final", "fila_excel_celular", "direccion_final", "fila_excel_direccion",
        "Cedula", "cc_limpia", "APELLIDOS Y NOMBRES DEL CLIENTE", "FECHA_RAD", "fecha_final", "criterio_fecha",
        "No.Lib", "libranza_limpia", "TIPO CREDITO", "MONTO", "monto_parseado", "MESES", "plazo_parseado",
        "estado_credito", "fecha_fin_estimada", "PAGADURIA", "COOPERATIVA",
    ]
    add_sheet("03_antes_despues", row_headers, prepared.rows)
    add_sheet("04_omitidos", row_headers, prepared.rejected)
    style_sheet(workbook["00_resumen"])
    workbook.save(workbook_path)


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        letter = column_cells[0].column_letter
        max_len = max((len(str(cell.value)) for cell in column_cells[:200] if cell.value is not None), default=8)
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)


def create_backup(output_dir: Path) -> Path:
    pg_dump = shutil.which("pg_dump")
    if not pg_dump:
        raise RuntimeError("pg_dump no esta disponible en este contenedor/entorno")
    output_dir.mkdir(parents=True, exist_ok=True)
    backup_path = output_dir / f"backup_pre_migracion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.sql"
    env = {**os.environ, "PGPASSWORD": settings.POSTGRES_PASSWORD}
    subprocess.run(
        [
            pg_dump,
            "-h",
            settings.POSTGRES_HOST,
            "-p",
            str(settings.POSTGRES_PORT),
            "-U",
            settings.POSTGRES_USER,
            "-d",
            settings.POSTGRES_DB,
            "-f",
            str(backup_path),
        ],
        check=True,
        env=env,
    )
    return backup_path


def apply_import(prepared: PreparedData, output_dir: Path, require_empty_db: bool) -> dict[str, Any]:
    db = SessionLocal()
    try:
        counts = db.execute(
            text(
                "SELECT "
                "(SELECT count(*) FROM creditos) AS creditos, "
                "(SELECT count(*) FROM pensionados) AS pensionados, "
                "(SELECT count(*) FROM cooperativas) AS cooperativas, "
                "(SELECT count(*) FROM pagadurias) AS pagadurias"
            )
        ).mappings().one()
        if not require_empty_db and any(counts[key] for key in counts.keys()):
            raise RuntimeError(f"La DB no esta vacia. Conteos actuales: {dict(counts)}. Usa --confirm-clean para limpiar y aplicar.")

        db.execute(
            text(
                """
                TRUNCATE TABLE
                    notificacion_lecturas,
                    notificaciones,
                    oportunidades_refinanciacion,
                    refinanciaciones,
                    pendientes_credito,
                    seguimientos,
                    documentos,
                    historial_creditos,
                    creditos,
                    pensionado_oficinas,
                    pensionados,
                    cooperativa_refinanciacion_reglas,
                    cooperativas,
                    pagadurias
                RESTART IDENTITY CASCADE
                """
            )
        )

        office_id = ensure_office(db)
        user_id = ensure_system_user(db, office_id)
        pagadurias: dict[str, int] = {}
        cooperativas: dict[str, int] = {}
        pensionados: dict[str, int] = {}

        for row in prepared.migratable:
            cc = row["cc_limpia"]
            if cc not in pensionados:
                final = next(item for item in prepared.pensionados if item["documento"] == cc)
                pensionados[cc] = db.execute(
                    text(
                        "INSERT INTO pensionados (oficina_id, created_by, nombre, segundo_nombre, apellidos, genero, documento, fecha_nacimiento, correo, telefono, celular, direccion, is_active, created_at, updated_at) "
                        "VALUES (:oficina_id, :created_by, :nombre, NULL, NULL, NULL, :documento, NULL, NULL, :telefono, :celular, :direccion, true, now(), now()) RETURNING id"
                    ),
                    {
                        "oficina_id": office_id,
                        "created_by": user_id,
                        "nombre": final["nombre"],
                        "documento": cc,
                        "telefono": final["telefono"] or None,
                        "celular": final["celular"] or None,
                        "direccion": final["direccion"] or None,
                    },
                ).scalar_one()
                db.execute(
                    text(
                        "INSERT INTO pensionado_oficinas (pensionado_id, oficina_id, created_by, vinculada_en, is_active, created_at, updated_at) "
                        "VALUES (:pensionado_id, :oficina_id, :created_by, now(), true, now(), now())"
                    ),
                    {"pensionado_id": pensionados[cc], "oficina_id": office_id, "created_by": user_id},
                )
            pagaduria_id = get_or_create_pagaduria(db, row["pagaduria_normalizada"], pagadurias)
            cooperativa_id = get_or_create_cooperativa(db, row["cooperativa_normalizada"], cooperativas)
            db.execute(
                text(
                    "INSERT INTO creditos (pensionado_id, asesor_id, oficina_id, cooperativa_id, pagaduria_id, nro_libranza, tipo_credito, monto_solicitado, monto_aprobado, plazo, estado, fecha_fin_estimada, observaciones, tiene_documentos_pendientes, documentos_pendientes, fecha_registro, is_active, created_at, updated_at) "
                    "VALUES (:pensionado_id, :asesor_id, :oficina_id, :cooperativa_id, :pagaduria_id, :nro_libranza, :tipo_credito, :monto, :monto, :plazo, :estado, :fecha_fin, NULL, false, NULL, :fecha_registro, true, now(), now())"
                ),
                {
                    "pensionado_id": pensionados[cc],
                    "asesor_id": user_id,
                    "oficina_id": office_id,
                    "cooperativa_id": cooperativa_id,
                    "pagaduria_id": pagaduria_id,
                    "nro_libranza": row["libranza_limpia"] or None,
                    "tipo_credito": row["tipo_normalizado"][:50] or None,
                    "monto": Decimal(row["monto_parseado"]),
                    "plazo": row["plazo_parseado"],
                    "estado": row["estado_credito"],
                    "fecha_fin": row["fecha_fin_estimada"],
                    "fecha_registro": datetime.combine(row["fecha_final"], datetime.min.time()),
                },
            )
        db.commit()
        return {"creditos_insertados": len(prepared.migratable), "pensionados_insertados": len(pensionados)}
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def ensure_office(db) -> int:
    office_id = db.execute(text("SELECT id FROM oficinas ORDER BY id LIMIT 1")).scalar()
    if office_id:
        return int(office_id)
    return int(
        db.execute(
            text("INSERT INTO oficinas (nombre, direccion, color, is_active, created_at, updated_at) VALUES ('Oficina Principal', 'Sin registrar', 'blue', true, now(), now()) RETURNING id")
        ).scalar_one()
    )


def ensure_system_user(db, office_id: int) -> int:
    user_id = db.execute(text("SELECT id FROM usuarios WHERE correo=:correo"), {"correo": SYSTEM_USER_EMAIL}).scalar()
    if user_id:
        db.execute(text("UPDATE usuarios SET oficina_id=:oficina_id, is_active=true WHERE id=:id"), {"oficina_id": office_id, "id": user_id})
        return int(user_id)
    return int(
        db.execute(
            text(
                "INSERT INTO usuarios (oficina_id, nombre, documento, correo, contrasena, rol, intentos_fallidos, is_active, created_at, updated_at) "
                "VALUES (:oficina_id, :nombre, :documento, :correo, :contrasena, 'asesora', 0, true, now(), now()) RETURNING id"
            ),
            {
                "oficina_id": office_id,
                "nombre": SYSTEM_USER_NAME,
                "documento": SYSTEM_USER_DOCUMENT,
                "correo": SYSTEM_USER_EMAIL,
                "contrasena": hash_password("Crediconfiemos2026*"),
            },
        ).scalar_one()
    )


def get_or_create_pagaduria(db, name: str, cache: dict[str, int]) -> int:
    if name in cache:
        return cache[name]
    cache[name] = int(db.execute(text("INSERT INTO pagadurias (nombre, is_active, created_at, updated_at) VALUES (:name, true, now(), now()) RETURNING id"), {"name": name}).scalar_one())
    return cache[name]


def get_or_create_cooperativa(db, name: str, cache: dict[str, int]) -> int:
    if name in cache:
        return cache[name]
    cache[name] = int(
        db.execute(
            text(
                "INSERT INTO cooperativas (nombre, edad_minima, edad_maxima, monto_minimo, monto_maximo, plazo_minimo, plazo_maximo, is_active, created_at, updated_at) "
                "VALUES (:name, 1, 130, 1, 9999999999, 1, 240, true, now(), now()) RETURNING id"
            ),
            {"name": name},
        ).scalar_one()
    )
    return cache[name]


def main() -> None:
    parser = argparse.ArgumentParser(description="Migracion controlada de creditos.xlsx")
    parser.add_argument("--input", default="/app/creditos.xlsx")
    parser.add_argument("--output-dir", default="/app/outputs/migracion_creditos")
    parser.add_argument("--dry-run", action="store_true", default=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--backup", action="store_true")
    parser.add_argument("--confirm-clean", action="store_true")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_dir = Path(args.output_dir)
    prepared = prepare(input_path)
    reports = write_reports(prepared, output_dir)
    result: dict[str, Any] = {"mode": "dry-run", "reports": reports}

    if args.apply:
        if not args.confirm_clean:
            raise SystemExit("Para aplicar debes pasar --confirm-clean. Dry-run generado, no se toco la DB.")
        if args.backup:
            result["backup"] = str(create_backup(output_dir))
        result["mode"] = "apply"
        result["apply"] = apply_import(prepared, output_dir, require_empty_db=True)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
