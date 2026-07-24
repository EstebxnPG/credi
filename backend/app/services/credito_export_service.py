from datetime import date
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.db.models.credito import Credito
from app.db.models.usuario import Usuario


DEFAULT_COLUMNS = [
    "tipo_credito",
    "fecha_registro",
    "nro_libranza",
    "pensionado",
    "monto",
    "meses",
    "cedula",
    "telefono",
    "celular",
    "pagaduria",
    "cooperativa",
    "cedula_asesor",
    "direccion",
    "barrio",
]

COLUMNS = {
    "tipo_credito": ("TIPO CREDITO", lambda c: c.tipo_credito),
    "fecha_registro": (
        "FECHA_RAD",
        lambda c: c.fecha_registro.replace(tzinfo=None) if c.fecha_registro else None,
    ),
    "nro_libranza": ("No.Lib", lambda c: c.nro_libranza),
    "pensionado": (
        "APELLIDOS Y NOMBRES DEL CLIENTE",
        lambda c: c.pensionado.nombre_completo,
    ),
    "monto": (
        "MONTO",
        lambda c: float(
            c.monto_aprobado if c.monto_aprobado is not None else c.monto_solicitado
        ),
    ),
    "meses": ("MESES", lambda c: c.plazo),
    "cedula": ("Cedula", lambda c: c.pensionado.documento),
    "telefono": ("TELEFONO", lambda c: c.pensionado.telefono),
    "correo": ("Correo", lambda c: c.pensionado.correo),
    "celular": ("CELULAR", lambda c: c.pensionado.celular),
    "pagaduria": ("PAGADURIA", lambda c: c.pagaduria.nombre),
    "cooperativa": ("COOPERATIVA", lambda c: c.cooperativa.nombre),
    "cedula_asesor": ("Cedula_Asesor", lambda c: c.asesor.documento),
    "direccion": ("DIRECCION", lambda c: c.pensionado.direccion),
    "barrio": ("BARRIO", lambda c: None),
    "credito_id": ("Credito", lambda c: c.id),
    "estado": ("Estado", lambda c: c.estado),
    "asesor": ("Asesor", lambda c: c.asesor.nombre),
    "oficina": ("Oficina", lambda c: c.oficina.nombre),
    "monto_solicitado": ("Monto solicitado", lambda c: float(c.monto_solicitado)),
    "monto_aprobado": (
        "Monto aprobado",
        lambda c: float(c.monto_aprobado) if c.monto_aprobado is not None else None,
    ),
    "plazo": ("Plazo", lambda c: c.plazo),
}


def exportar_creditos(
    db: Session,
    usuario: Usuario,
    columnas: list[str],
    desde: date | None,
    hasta: date | None,
    oficina_id: int | None,
    monto_desde: float | None,
    monto_hasta: float | None,
) -> BytesIO:
    invalidas = [columna for columna in columnas if columna not in COLUMNS]
    if invalidas or not columnas:
        raise HTTPException(
            status_code=422,
            detail=f"Columnas invalidas: {', '.join(invalidas) or 'ninguna seleccionada'}",
        )
    if desde and hasta and desde > hasta:
        raise HTTPException(
            status_code=422,
            detail="La fecha desde no puede ser posterior a la fecha hasta",
        )
    if monto_desde is not None and monto_hasta is not None and monto_desde > monto_hasta:
        raise HTTPException(
            status_code=422,
            detail="El monto minimo no puede superar el monto maximo",
        )

    query = (
        db.query(Credito)
        .options(
            joinedload(Credito.pensionado),
            joinedload(Credito.pagaduria),
            joinedload(Credito.cooperativa),
            joinedload(Credito.asesor),
            joinedload(Credito.oficina),
        )
        .filter(Credito.is_active == True)  # noqa: E712
    )
    if usuario.rol != "administrador":
        query = query.filter(Credito.oficina_id == usuario.oficina_id)
    elif oficina_id:
        query = query.filter(Credito.oficina_id == oficina_id)
    if desde:
        query = query.filter(Credito.fecha_registro >= desde)
    if hasta:
        query = query.filter(Credito.fecha_registro < date.fromordinal(hasta.toordinal() + 1))

    creditos = query.order_by(Credito.fecha_registro.asc(), Credito.id.asc()).all()

    def monto(credito: Credito) -> float:
        return float(
            credito.monto_aprobado
            if credito.monto_aprobado is not None
            else credito.monto_solicitado
        )

    creditos = [
        credito
        for credito in creditos
        if (monto_desde is None or monto(credito) >= monto_desde)
        and (monto_hasta is None or monto(credito) <= monto_hasta)
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Creditos"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.append([COLUMNS[columna][0] for columna in columnas])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")

    def excel_safe(value):
        return (
            f"'{value}"
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
            else value
        )

    for credito in creditos:
        ws.append([excel_safe(COLUMNS[columna][1](credito)) for columna in columnas])

    ws.auto_filter.ref = ws.dimensions
    for index, key in enumerate(columnas, 1):
        letter = get_column_letter(index)
        values = [
            str(ws.cell(row, index).value or "")
            for row in range(1, min(ws.max_row, 200) + 1)
        ]
        ws.column_dimensions[letter].width = min(
            max(max(map(len, values), default=10) + 2, 12),
            36,
        )
        if key == "fecha_registro":
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm"
        if key in {"monto", "monto_solicitado", "monto_aprobado"}:
            for cell in ws[letter][1:]:
                cell.number_format = "$ #,##0.00"

    ws.row_dimensions[1].height = 24
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
